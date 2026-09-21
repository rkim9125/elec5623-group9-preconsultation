import io
import sys
import zipfile

from docx import Document
from PIL import Image
from pypdf import PdfReader, PdfWriter
import pytest
from reportlab.pdfgen import canvas

from app.utils.document_errors import DocumentError
from app.utils import document_processing as processing
from app.utils.document_processing import extract_document, validate_file


def pdf_bytes(pages: list[str | None]) -> bytes:
    stream = io.BytesIO()
    pdf = canvas.Canvas(stream)
    for text in pages:
        if text:
            pdf.drawString(40, 700, text)
        pdf.showPage()
    pdf.save()
    return stream.getvalue()


def image_bytes(format="PNG", size=(20, 20)) -> bytes:
    stream = io.BytesIO()
    Image.new("RGB", size, "white").save(stream, format=format)
    return stream.getvalue()


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Patient reports 2 days of cough.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Test"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Hb"
    table.cell(1, 1).text = "11.2 g/dL"
    document.add_paragraph("No medication reported.")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def assert_code(code, action):
    with pytest.raises(DocumentError) as error:
        action()
    assert error.value.code == code


def test_utf8_lines_retain_unicode_and_locations():
    result = extract_document("notes.TXT", "\ufeffNo allergy reported.\n\n症状持续两天".encode())
    assert result.status == "processed"
    assert [(chunk.text, chunk.location, chunk.method) for chunk in result.chunks] == [
        ("No allergy reported.", {"line": 1}, "text"),
        ("症状持续两天", {"line": 3}, "text"),
    ]


def test_pdf_returns_one_based_page_chunks():
    result = extract_document("report.pdf", pdf_bytes(["Hb: 11.2 g/dL", "Date: 2026-09-22"]))
    assert result.status == "processed"
    assert [chunk.location for chunk in result.chunks] == [{"page": 1}, {"page": 2}]
    assert "11.2 g/dL" in result.chunks[0].text


def test_docx_body_order_and_table_locations():
    result = extract_document("report.docx", docx_bytes())
    assert result.status == "processed"
    assert [chunk.text for chunk in result.chunks] == [
        "Patient reports 2 days of cough.", "Test", "Value", "Hb", "11.2 g/dL", "No medication reported."
    ]
    assert result.chunks[4].location["table"] == 1
    assert result.chunks[4].location["row"] == 2
    assert result.chunks[4].location["column"] == 2


@pytest.mark.parametrize("name", ["../a.txt", "/a.txt", "a\\b.txt", "a\x00.txt", "a\n.txt", " a.txt", "a\u202e.txt", "x" * 260 + ".txt"])
def test_unsafe_filename_rejected(name):
    assert_code("INVALID_FILENAME", lambda: validate_file(name, b"text", None, 1000))


@pytest.mark.parametrize("name,data,mime,code", [
    ("a.doc", b"legacy", None, "UNSUPPORTED_FILE_TYPE"),
    ("a.pdf", b"not PDF", None, "UNSUPPORTED_FILE_TYPE"),
    ("a.png", image_bytes("JPEG"), None, "UNSUPPORTED_FILE_TYPE"),
    ("a.jpg", image_bytes(), None, "UNSUPPORTED_FILE_TYPE"),
    ("a.docx", b"not a zip", None, "UNSUPPORTED_FILE_TYPE"),
    ("a.txt", b"text", "application/pdf", "UNSUPPORTED_FILE_TYPE"),
    ("a.txt", b"", None, "EMPTY_FILE"),
    ("a.txt", b"\xff\xfe", None, "DOCUMENT_PARSE_FAILED"),
    ("a.txt", b"hello\x00world", None, "DOCUMENT_PARSE_FAILED"),
])
def test_file_validation_errors(name, data, mime, code):
    assert_code(code, lambda: validate_file(name, data, mime, 10000))


def test_generic_mime_and_utf8_charset_accepted():
    assert validate_file("A.TXT", b"text", "text/plain; charset=utf-8", 100) == ("A.TXT", "text/plain")
    assert validate_file("a.pdf", pdf_bytes(["text"]), "application/octet-stream", 10000)[1] == "application/pdf"


def test_limits_do_not_truncate_source():
    assert_code("FILE_TOO_LARGE", lambda: validate_file("a.txt", b"12345", None, 4))
    assert_code("DOCUMENT_LIMIT_EXCEEDED", lambda: extract_document("a.txt", b"12345", max_text_chars=4))
    assert_code("DOCUMENT_LIMIT_EXCEEDED", lambda: extract_document("a.pdf", pdf_bytes(["1", "2"]), max_pdf_pages=1))
    assert_code("DOCUMENT_LIMIT_EXCEEDED", lambda: extract_document("a.png", image_bytes(), max_image_pixels=100))
    assert_code("DOCUMENT_LIMIT_EXCEEDED", lambda: extract_document("a.docx", docx_bytes(), max_text_chars=10))


@pytest.mark.parametrize("filename,data", [
    ("a.pdf", b"%PDF-1.4\nbroken"),
    ("a.docx", b"PK\x03\x04broken"),
    ("a.png", b"\x89PNG\r\n\x1a\nbroken"),
])
def test_malformed_supported_formats_are_safe_errors(filename, data):
    assert_code("DOCUMENT_PARSE_FAILED", lambda: extract_document(filename, data))


def test_encrypted_pdf_is_rejected():
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(pdf_bytes(["confidential"]))))
    writer.encrypt("secret")
    stream = io.BytesIO()
    writer.write(stream)
    assert_code("DOCUMENT_PARSE_FAILED", lambda: extract_document("encrypted.pdf", stream.getvalue()))


def test_docx_rejects_unsafe_archives_and_expanded_size(monkeypatch):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("../payload", "text")
    assert_code("DOCUMENT_PARSE_FAILED", lambda: extract_document("unsafe.docx", stream.getvalue()))
    monkeypatch.setattr(processing, "MAX_DOCX_UNCOMPRESSED_BYTES", 100)
    assert_code("DOCUMENT_LIMIT_EXCEEDED", lambda: extract_document("large.docx", docx_bytes()))


@pytest.mark.parametrize("name,data", [("scan.png", image_bytes()), ("scan.jpg", image_bytes("JPEG")), ("scan.pdf", pdf_bytes([None]))])
def test_ocr_disabled_reports_no_success(name, data, monkeypatch):
    def unexpected(*args):
        pytest.fail("OCR must not execute when disabled")
    monkeypatch.setattr(processing, "_ocr_image", unexpected)
    result = extract_document(name, data)
    assert result.status == "ocr_required"
    assert result.chunks == []
    assert result.warnings[0].startswith("OCR_DISABLED:")


def test_mixed_pdf_reports_partial_and_preserves_text():
    result = extract_document("mixed.pdf", pdf_bytes(["Recognized text", None]))
    assert result.status == "partial"
    assert len(result.chunks) == 1
    assert "page 2" in result.warnings[0]


def test_enabled_image_and_pdf_ocr(monkeypatch):
    calls = []
    def ocr(image, language, timeout):
        calls.append((image.size, language, timeout))
        return "Glucose: 5.4 mmol/L"
    monkeypatch.setattr(processing, "_ocr_image", ocr)
    monkeypatch.setattr(processing, "_ocr_pdf_page", lambda *args: "Glucose: 5.4 mmol/L")
    image_result = extract_document("scan.png", image_bytes(), ocr_enabled=True, ocr_language="eng+chi_sim", ocr_timeout_seconds=3)
    pdf_result = extract_document("scan.pdf", pdf_bytes([None]), ocr_enabled=True)
    for result in (image_result, pdf_result):
        assert result.status == "processed"
        assert result.chunks[0].method == "ocr"
        assert result.chunks[0].text == "Glucose: 5.4 mmol/L"
        assert result.warnings[0].startswith("OCR_REVIEW_REQUIRED:")
    assert calls[0] == ((20, 20), "eng+chi_sim", 3)


def test_optional_pdf_rasterizer_feeds_page_to_ocr(monkeypatch):
    pytest.importorskip("pypdfium2")
    sizes = []
    def ocr(image, *args):
        sizes.append(image.size)
        return "Recognized scan"
    monkeypatch.setattr(processing, "_ocr_image", ocr)
    result = extract_document("scan.pdf", pdf_bytes([None]), ocr_enabled=True)
    assert result.status == "processed"
    assert result.chunks[0].text == "Recognized scan"
    assert sizes[0][0] > 20


@pytest.mark.parametrize("module,filename,data", [
    ("pytesseract", "scan.png", image_bytes()),
    ("pypdfium2", "scan.pdf", pdf_bytes([None])),
])
def test_missing_optional_ocr_dependencies_return_actionable_status(module, filename, data, monkeypatch):
    monkeypatch.setitem(sys.modules, module, None)
    result = extract_document(filename, data, ocr_enabled=True)
    assert result.status == "ocr_required"
    assert "OCR_UNAVAILABLE" in result.warnings[0]


@pytest.mark.parametrize("code", ["OCR_UNAVAILABLE", "OCR_TIMEOUT"])
def test_ocr_operational_failure_is_visible(code, monkeypatch):
    def unavailable(*args):
        raise DocumentError(code, "test failure")
    monkeypatch.setattr(processing, "_ocr_image", unavailable)
    result = extract_document("scan.png", image_bytes(), ocr_enabled=True)
    assert result.status == "ocr_required"
    assert result.chunks == []
    assert code in result.warnings[0]


def test_empty_ocr_and_empty_text_are_distinguished(monkeypatch):
    monkeypatch.setattr(processing, "_ocr_image", lambda *args: "  ")
    result = extract_document("scan.png", image_bytes(), ocr_enabled=True)
    assert result.status == "ocr_required"
    assert "OCR_NO_TEXT" in result.warnings[0]
    assert_code("EMPTY_DOCUMENT", lambda: extract_document("empty.txt", b" \n"))


def test_docx_images_are_explicitly_incomplete():
    document = Document()
    document.add_paragraph("See attached image")
    document.add_picture(io.BytesIO(image_bytes()))
    stream = io.BytesIO()
    document.save(stream)
    result = extract_document("with-image.docx", stream.getvalue())
    assert result.status == "partial"
    assert "DOCX_IMAGES_NOT_EXTRACTED" in result.warnings[0]
