import io
import re
import zipfile

from docx import Document
from pypdf import PdfReader
import pytest

from app.utils.document_errors import DocumentError
from app.utils.document_rendering import render_summary


FACTS = {
    "Symptoms": "Cough for 2 days. Temperature 38.2 C. No shortness of breath reported.",
    "History": "Patient says: <b>no allergy</b> & no medicine. Hb < 12 g/dL.",
    "Patient statement": "症状持续两天，未报告过敏。",
}
QUESTIONS = ["When should I seek help?", "Does Hb 11.2 g/dL need follow-up?"]


def render(format, sections=None, questions=None, **kwargs):
    return render_summary(
        FACTS if sections is None else sections,
        QUESTIONS if questions is None else questions,
        session_id="session-42", summary_ref="summary-v3", format=format, **kwargs,
    )


@pytest.mark.parametrize("format", ["pdf", "docx"])
def test_roundtrip_retains_facts_order_questions_and_metadata(format):
    content = render(format)
    if format == "pdf":
        document = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() for page in document.pages)
    else:
        document = Document(io.BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for literal in [*FACTS.keys(), *FACTS.values(), *QUESTIONS, "session-42", "summary-v3"]:
        # PDF wraps lines but does not change the facts or turn source markup into styling.
        assert " ".join(literal.split()) in " ".join(text.split())
    assert text.index("Symptoms") < text.index("History") < text.index("Patient statement")
    assert "<b>no allergy</b>" in text


def test_reportlab_markup_cannot_create_images_or_links():
    literal = '<img src="https://example.invalid/image.png"/> <a href="file:///etc/passwd">literal</a>'
    content = render("pdf", sections={"Untrusted text": literal})
    reader = PdfReader(io.BytesIO(content))
    text = " ".join("\n".join(page.extract_text() for page in reader.pages).split())
    assert literal in text
    assert not any(page.get("/Annots") for page in reader.pages)


def test_docx_does_not_create_external_relationships_from_source_markup():
    literal = '<a href="https://example.invalid">no allergy</a>'
    content = render("docx", sections={"Source": literal})
    document = Document(io.BytesIO(content))
    assert literal in [paragraph.text for paragraph in document.paragraphs]
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert b'TargetMode="External"' not in archive.read("word/_rels/document.xml.rels")


def test_long_pdf_sections_paginate_without_losing_final_facts():
    # A single long paragraph must split across pages instead of raising LayoutError.
    value = "Patient reports stable symptoms; no new medication. " * 900 + "FINAL FACT: Hb 11.2 g/dL."
    content = render("pdf", sections={"Long history": value})
    reader = PdfReader(io.BytesIO(content))
    assert len(reader.pages) >= 5
    # Running page numbers can appear between the two halves of a wrapped sentence.
    extracted = " ".join(re.sub(r"Page \d+", "", " ".join(page.extract_text() for page in reader.pages)).split())
    assert "FINAL FACT: Hb 11.2 g/dL." in extracted
    assert extracted.count("Patient reports stable symptoms;") == 900
    assert "Patient reports stable symptoms;" in reader.pages[0].extract_text()


def test_long_docx_retains_every_paragraph_and_questions():
    lines = [f"Observation {index}: value 11.2 g/dL." for index in range(500)]
    document = Document(io.BytesIO(render("docx", sections={"Observations": "\n".join(lines)})))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert all(line in paragraphs for line in lines)
    assert f"2. {QUESTIONS[1]}" in paragraphs


def test_pdf_unknown_glyph_fails_instead_of_silent_replacement(monkeypatch):
    monkeypatch.delenv("C5_PDF_FONT_PATH", raising=False)
    with pytest.raises(DocumentError) as error:
        render("pdf", sections={"Source": "Patient statement 🙂"})
    assert error.value.code == "DOCUMENT_RENDER_FAILED"
    assert "C5_PDF_FONT_PATH" in error.value.message


def test_font_misconfiguration_has_stable_error(monkeypatch):
    monkeypatch.setenv("C5_PDF_FONT_PATH", "/nonexistent/c5-font.ttf")
    with pytest.raises(DocumentError) as error:
        render("pdf")
    assert error.value.code == "DOCUMENT_RENDER_FAILED"


def test_explicit_font_path_overrides_environment(monkeypatch):
    monkeypatch.setenv("C5_PDF_FONT_PATH", "/nonexistent/environment-font.ttf")
    # An explicit empty path requests the built-in fonts regardless of environment.
    assert render("pdf", font_path="").startswith(b"%PDF")
    with pytest.raises(DocumentError) as error:
        render("pdf", font_path="/nonexistent/explicit-font.ttf")
    assert error.value.code == "DOCUMENT_RENDER_FAILED"


@pytest.mark.parametrize("format", ["pdf", "docx"])
def test_exports_reject_oversize_and_control_characters(format):
    with pytest.raises(DocumentError) as error:
        render(format, max_text_chars=10)
    assert error.value.code == "DOCUMENT_LIMIT_EXCEEDED"
    with pytest.raises(DocumentError) as error:
        render(format, sections={"Source": "value\x00hidden"})
    assert error.value.code == "DOCUMENT_RENDER_FAILED"


def test_docx_can_preserve_unicode_outside_default_pdf_font_repertoire():
    literal = "Patient statement: العربية हिन्दी 🙂"
    document = Document(io.BytesIO(render("docx", sections={"Source": literal})))
    assert literal in [paragraph.text for paragraph in document.paragraphs]


def test_export_unsupported_format_rejected():
    with pytest.raises(DocumentError) as error:
        render("html")
    assert error.value.code == "UNSUPPORTED_FILE_TYPE"
