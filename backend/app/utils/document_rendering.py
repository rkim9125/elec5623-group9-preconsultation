"""Faithful, paginated exports of already approved C3 summaries (no LLM calls)."""

from __future__ import annotations

from hashlib import sha256
import io
import os
from pathlib import Path
from threading import RLock
from typing import Literal
import unicodedata
from xml.sax.saxutils import escape

from app.utils.document_errors import DocumentError


_FONT_LOCK = RLock()


def _pdf_font_selector(font_path: str | None = None):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont

    configured = os.environ.get("C5_PDF_FONT_PATH", "") if font_path is None else font_path
    with _FONT_LOCK:
        if configured:
            path = Path(configured).expanduser()
            if not path.is_file():
                raise DocumentError("DOCUMENT_RENDER_FAILED", "C5_PDF_FONT_PATH does not identify a readable TTF font.")
            name = "C5Custom" + sha256(str(path.resolve()).encode()).hexdigest()[:12]
            try:
                if name not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(name, str(path)))
                supported = pdfmetrics.getFont(name).face.charToGlyph
            except Exception as exc:
                raise DocumentError("DOCUMENT_RENDER_FAILED", "C5_PDF_FONT_PATH must be a supported TrueType font.") from exc

            def select(char: str) -> str:
                if ord(char) not in supported:
                    raise DocumentError("DOCUMENT_RENDER_FAILED", f"Configured PDF font lacks U+{ord(char):04X}; configure a font covering the summary language.")
                return name

            return select
        if "STSong-Light" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))

    def select(char: str) -> str:
        try:
            char.encode("cp1252")
            return "Helvetica"
        except UnicodeEncodeError:
            pass
        try:
            # Restrict the CID font to its supported Simplified Chinese repertoire.
            # Do not silently emit missing glyphs for emoji or arbitrary scripts.
            char.encode("gb2312")
            return "STSong-Light"
        except UnicodeEncodeError as exc:
            raise DocumentError(
                "DOCUMENT_RENDER_FAILED",
                f"Default PDF fonts do not support U+{ord(char):04X}; set C5_PDF_FONT_PATH to a TTF covering this language, or export DOCX.",
            ) from exc

    return select


def _pdf_markup(text: str, select_font) -> str:
    """Escape all caller markup before applying trusted per-font runs."""
    runs: list[str] = []
    current_font: str | None = None
    current_text: list[str] = []

    def flush() -> None:
        if current_text:
            literal = escape("".join(current_text)).replace("\t", "&#160;" * 4)
            runs.append(f'<font name="{current_font}">{literal}</font>')
            current_text.clear()

    for char in text:
        if char in "\n\r":
            flush()
            # CRLF was normalized by the caller; each newline is retained.
            runs.append("<br/>")
            continue
        font = select_font(" " if char == "\t" else char)
        if font != current_font:
            flush()
            current_font = font
        current_text.append(char)
    flush()
    return "".join(runs) or "&#160;"


def _render_pdf(sections, questions, session_id, summary_ref, font_path: str | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import CondPageBreak, Paragraph, SimpleDocTemplate, Spacer

    output = io.BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm,
        topMargin=18 * mm, bottomMargin=19 * mm, title="Preconsultation summary",
        author="", subject="Approved preconsultation summary",
    )
    select_font = _pdf_font_selector(font_path)
    body = ParagraphStyle(
        "C5Body", fontName="Helvetica", fontSize=10, leading=15,
        spaceAfter=9, alignment=TA_LEFT, splitLongWords=True,
        allowWidows=0, allowOrphans=0,
    )
    title = ParagraphStyle("C5Title", parent=body, fontSize=18, leading=23, spaceAfter=13)
    heading = ParagraphStyle("C5Heading", parent=body, fontSize=12, leading=17, spaceBefore=10, spaceAfter=5)
    metadata = ParagraphStyle("C5Metadata", parent=body, fontSize=8, leading=12, textColor=colors.HexColor("#444444"))

    def paragraph(text, style=body):
        return Paragraph(_pdf_markup(text.replace("\r\n", "\n").replace("\r", "\n"), select_font), style)

    story = [paragraph("Preconsultation summary", title)]
    story += [paragraph(f"Session ID: {session_id}", metadata), paragraph(f"Summary reference: {summary_ref}", metadata), Spacer(1, 4 * mm)]
    for name, text in sections.items():
        # Reserve a heading plus two body lines without keeping an arbitrarily
        # large following paragraph together (which produces an empty title page).
        story.append(CondPageBreak(62))
        story.append(paragraph(name, heading))
        # Separate source paragraphs permit arbitrary-length sections to paginate.
        for line in text.split("\n"):
            story.append(paragraph(line))
    story.append(CondPageBreak(62))
    story.append(paragraph("Patient questions", heading))
    for index, question in enumerate(questions, 1):
        story.append(paragraph(f"{index}. {question}"))
    if not questions:
        story.append(paragraph("None recorded."))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawRightString(A4[0] - 20 * mm, 11 * mm, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _render_docx(sections, questions, session_id, summary_ref) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for style_name in ("Title", "Heading 1"):
        style = document.styles[style_name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        # Some python-docx templates include a blue rule in Title by default.
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    document.core_properties.title = "Preconsultation summary"
    document.core_properties.author = ""
    document.add_paragraph("Preconsultation summary", style="Title")
    document.add_paragraph(f"Session ID: {session_id}")
    document.add_paragraph(f"Summary reference: {summary_ref}")
    for name, text in sections.items():
        document.add_heading(name, level=1)
        for line in text.split("\n"):
            document.add_paragraph(line)
    document.add_heading("Patient questions", level=1)
    for index, question in enumerate(questions, 1):
        document.add_paragraph(f"{index}. {question}")
    if not questions:
        document.add_paragraph("None recorded.")
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def render_summary(
    sections: dict[str, str],
    patient_questions: list[str],
    *,
    session_id: str,
    summary_ref: str,
    format: Literal["pdf", "docx"],
    max_text_chars: int = 200000,
    font_path: str | None = None,
) -> bytes:
    """Render literal ordered text and metadata; no inference or state mutation.

    PDF defaults support Western Latin and GB2312 Chinese. Set C5_PDF_FONT_PATH
    to an installed TrueType font for other scripts; missing glyphs raise an
    explicit error. DOCX retains Unicode and relies on the viewer's font fallback.
    """
    if format not in {"pdf", "docx"}:
        raise DocumentError("UNSUPPORTED_FILE_TYPE", "Summary exports support PDF and DOCX.")
    values = [session_id, summary_ref, *sections.keys(), *sections.values(), *patient_questions]
    if any(not isinstance(value, str) for value in values):
        raise DocumentError("DOCUMENT_RENDER_FAILED", "All summary fields must be text.")
    if max_text_chars <= 0 or sum(map(len, values)) > max_text_chars:
        raise DocumentError("DOCUMENT_LIMIT_EXCEEDED", "Summary exceeds the export character limit.")
    if any(unicodedata.category(char) == "Cc" and char not in "\t\n\r" for value in values for char in value):
        raise DocumentError("DOCUMENT_RENDER_FAILED", "Summary contains unsupported control characters.")
    try:
        if format == "pdf":
            return _render_pdf(sections, patient_questions, session_id, summary_ref, font_path)
        return _render_docx(sections, patient_questions, session_id, summary_ref)
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("DOCUMENT_RENDER_FAILED", "Summary could not be rendered.") from exc
